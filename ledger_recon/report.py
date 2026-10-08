"""Excel report building and styling."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

DEFAULT_HEADER_COLOR = "305496"


def style_workbook(path: Path, colors: dict | None = None, default: str = DEFAULT_HEADER_COLOR) -> None:
    """Bold white headers on a coloured fill, frozen header row, sensible column widths."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    colors = colors or {}
    wb = load_workbook(path)
    font = Font(bold=True, color="FFFFFF", name="Arial")
    for ws in wb.worksheets:
        fill = PatternFill("solid", start_color=colors.get(ws.title, default))
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font, cell.fill = font, fill
        for column in ws.columns:
            longest = max((len(str(c.value)) if c.value is not None else 0) for c in column[:200])
            ws.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 10), 45)
    wb.save(path)


def account_pivot(statement: pd.DataFrame) -> pd.DataFrame:
    if statement.empty:
        return pd.DataFrame(columns=["account_no", "posted_date", "direction", "rows", "amount"])
    return (
        statement.groupby(["account_no", "posted_date", "direction"])
        .agg(rows=("amount", "size"), amount=("amount", "sum"))
        .reset_index()
    )


def summary_frame(summary: dict, extra: dict | None = None) -> pd.DataFrame:
    rows = {**summary, **(extra or {})}
    return pd.DataFrame({"metric": list(rows), "value": list(rows.values())})


def write_report(sheets: dict, out_path: Path) -> None:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    style_workbook(out_path)
